from pathlib import Path
from io import BytesIO
from config import TEXT_EXTENSIONS

def extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        try: reader.decrypt("")
        except Exception: return ""
    pages = [page.extract_text().strip() for page in reader.pages if page.extract_text()]
    return "\n\n".join(pages).strip()

def extract_docx(data: bytes) -> str:
    import docx
    document = docx.Document(BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells: parts.append(" | ".join(cells))
    return "\n".join(parts).strip()

def extract_html(data: bytes) -> str:
    from bs4 import BeautifulSoup
    try: soup = BeautifulSoup(data, "lxml")
    except Exception: soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]): tag.decompose()
    return soup.get_text(separator="\n", strip=True)

def extract_epub(data: bytes) -> str:
    import tempfile, ebooklib, os
    from ebooklib import epub
    from bs4 import BeautifulSoup
    with tempfile.NamedTemporaryFile(suffix=".epub", delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        try: book = epub.read_epub(tmp_path, options={"ignore_ncx": True})
        except TypeError: book = epub.read_epub(tmp_path)
    finally:
        try: os.unlink(tmp_path)
        except OSError: pass
    parts = []
    for item in book.get_items():
        if item.get_type() == ebooklib.ITEM_DOCUMENT:
            try:
                soup = BeautifulSoup(item.get_content(), "html.parser")
                for tag in soup(["script", "style"]): tag.decompose()
                text = soup.get_text(separator="\n", strip=True)
                if text: parts.append(text)
            except Exception: continue
    return "\n\n".join(parts).strip()

def extract_text_from_file(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in TEXT_EXTENSIONS: return data.decode("utf-8", errors="ignore")
    if ext == ".pdf": return extract_pdf(data)
    if ext == ".docx": return extract_docx(data)
    if ext in {".html", ".htm"}: return extract_html(data)
    if ext == ".epub": return extract_epub(data)
    raise ValueError(f"Unsupported file extension: {ext}")